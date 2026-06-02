import os
import io
import threading
import time
from collections import Counter
from datetime import datetime, timedelta

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from database import db
from services.notificaciones import correo_habilitado_para, enviar_correo_con_adjunto
from services.reportes import generar_csv_logs


_scheduler_started = False
_scheduler_lock = threading.Lock()


def _parse_destinatarios(destinatarios_json):
    if not destinatarios_json:
        return []
    texto = str(destinatarios_json).strip()
    if not texto:
        return []
    if texto.startswith('['):
        try:
            import json
            datos = json.loads(texto)
            return [str(item).strip() for item in datos if str(item).strip()]
        except Exception:
            pass
    texto = texto.replace(';', ',')
    return [parte.strip() for parte in texto.split(',') if parte.strip()]


def _obtener_destinatarios_finales(reporte):
    destinatarios = _parse_destinatarios(reporte['destinatarios_json'])
    if destinatarios:
        return destinatarios
    fallback = os.environ.get('ALERT_EMAIL_TO', '')
    destinatarios = [x.strip() for x in fallback.replace(';', ',').split(',') if x.strip()]
    if destinatarios:
        return destinatarios
    return db.obtener_correos_usuarios()


def _construir_resumen_reporte(reporte):
    dias = int(reporte['frecuencia_dias'])
    fecha_inicio_dt = datetime.now() - timedelta(days=dias)
    fecha_inicio = fecha_inicio_dt.strftime('%Y-%m-%d %H:%M:%S')
    logs = db.obtener_logs_desde(fecha_inicio)
    logs_csv = [(fila['usuario'], fila['accion'], fila['detalles'], fila['fecha']) for fila in logs]

    inicios_sesion = [fila for fila in logs if 'inicio de ses' in str(fila['accion']).lower()]
    usuarios_ingreso = Counter(fila['usuario'] for fila in inicios_sesion)
    top_usuarios = usuarios_ingreso.most_common(5)
    estadisticas_actividades = db.obtener_todas_las_actividades_para_kpis()
    total_actividades = estadisticas_actividades.get('total', 0) or 0
    completadas = estadisticas_actividades.get('completadas', 0) or 0
    avance = (completadas / total_actividades * 100) if total_actividades else 0

    lineas_top = [f"- {usuario}: {conteo}" for usuario, conteo in top_usuarios] or ['- Sin ingresos registrados']
    cuerpo = [
        f"Reporte programado: {reporte['nombre']}",
        f"Periodo revisado: últimos {dias} días",
        '',
        'Ingresos al sistema:',
        *lineas_top,
        '',
        'Avance general:',
        f"- Total de actividades: {total_actividades}",
        f"- Completadas: {completadas}",
        f"- Pendientes: {estadisticas_actividades.get('pendientes', 0) or 0}",
        f"- En progreso: {estadisticas_actividades.get('en_progreso', 0) or 0}",
        f"- Avance estimado: {avance:.2f}%",
        '',
        'Detalle adjunto en CSV con los eventos del periodo.'
    ]

    asunto = f"Reporte programado de actividad - últimos {dias} días"

    if str(reporte.get('formato', 'csv')).lower() == 'word':
        doc = Document()
        titulo = doc.add_heading('Reporte Programado de Actividad', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Reporte: {reporte['nombre']}")
        doc.add_paragraph(f"Periodo revisado: últimos {dias} días")
        doc.add_heading('Ingresos al sistema', level=1)
        if top_usuarios:
            for usuario, conteo in top_usuarios:
                doc.add_paragraph(f'{usuario}: {conteo}', style='List Bullet')
        else:
            doc.add_paragraph('Sin ingresos registrados', style='List Bullet')
        doc.add_heading('Avance general', level=1)
        doc.add_paragraph(f"Total de actividades: {total_actividades}")
        doc.add_paragraph(f"Completadas: {completadas}")
        doc.add_paragraph(f"Pendientes: {estadisticas_actividades.get('pendientes', 0) or 0}")
        doc.add_paragraph(f"En progreso: {estadisticas_actividades.get('en_progreso', 0) or 0}")
        doc.add_paragraph(f"Avance estimado: {avance:.2f}%")
        doc.add_heading('Eventos del periodo', level=1)
        for fila in logs[:20]:
            doc.add_paragraph(f"{fila['fecha']} | {fila['usuario']} | {fila['accion']} | {fila['detalles']}", style='List Bullet')
        adjunto = io.BytesIO()
        doc.save(adjunto)
        adjunto.seek(0)
        nombre_adjunto = f"reporte_actividad_{dias}d_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    else:
        adjunto = generar_csv_logs(logs_csv)
        nombre_adjunto = f"reporte_actividad_{dias}d_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        mimetype = 'text/csv'

    return asunto, '\n'.join(cuerpo), adjunto, nombre_adjunto, mimetype, _obtener_destinatarios_finales(reporte)


def _procesar_reportes_programados():
    reportes = db.obtener_reportes_programados_vencidos(limite=25)
    for reporte in reportes:
        try:
            destinatarios = _obtener_destinatarios_finales(reporte)
            if not correo_habilitado_para(destinatarios):
                continue

            asunto, cuerpo, adjunto, nombre_adjunto, mimetype, destinatarios = _construir_resumen_reporte(reporte)
            envio = enviar_correo_con_adjunto(asunto, cuerpo, destinatarios, adjunto, nombre_adjunto, mimetype)
            if envio.get('ok'):
                ahora = datetime.now()
                proximo_envio = (ahora + timedelta(days=int(reporte['frecuencia_dias']))).strftime('%Y-%m-%d %H:%M:%S')
                db.marcar_reporte_programado_enviado(reporte['id'], ultimo_envio=ahora.strftime('%Y-%m-%d %H:%M:%S'), proximo_envio=proximo_envio)
                db.registrar_log('sistema', 'Reporte programado enviado', f"Reporte {reporte['id']} ({reporte['nombre']}) a {len(destinatarios)} destinatarios")
            else:
                db.registrar_log('sistema', 'Reporte programado fallido', f"Reporte {reporte['id']} ({reporte['nombre']}): {envio.get('error', envio.get('motivo', 'desconocido'))}")
        except Exception as exc:
            db.registrar_log('sistema', 'Reporte programado error', f"Reporte {reporte['id']}: {exc}")


def _loop_scheduler():
    intervalo = int(os.environ.get('REPORT_SCHEDULER_INTERVAL_SECONDS', '300'))
    while True:
        try:
            _procesar_reportes_programados()
        except Exception:
            pass
        time.sleep(intervalo)


def iniciar_scheduler_reportes_programados():
    global _scheduler_started
    if os.environ.get('ENABLE_REPORT_SCHEDULER', '1') == '0':
        return
    with _scheduler_lock:
        if _scheduler_started:
            return
        hilo = threading.Thread(target=_loop_scheduler, daemon=True, name='scheduled-report-scheduler')
        hilo.start()
        _scheduler_started = True
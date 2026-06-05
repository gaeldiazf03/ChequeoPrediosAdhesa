import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _estado_tarea(tarea):
	estado = str(tarea.get('estado', '') or '').strip().lower()
	if estado in ('verde', 'amarillo', 'rojo'):
		return estado
	if estado in ('completada', 'completado'):
		return 'verde'
	if estado in ('en_progreso', 'en progreso', 'proceso'):
		return 'amarillo'
	if tarea.get('completada'):
		return 'verde'
	return 'rojo'


def _a_float(valor, default=0.0):
	try:
		return float(valor or 0)
	except Exception:
		return default


def _agrupar_actividades_por_estado(tareas):
	agrupadas = {
		'completadas': [],
		'en_proceso': [],
		'no_iniciadas': []
	}

	for tarea in tareas or []:
		nombre = str(tarea.get('texto', 'Sin descripcion') or 'Sin descripcion')
		estado = _estado_tarea(tarea)
		if estado == 'verde':
			agrupadas['completadas'].append(nombre)
		elif estado == 'amarillo':
			agrupadas['en_proceso'].append(nombre)
		else:
			agrupadas['no_iniciadas'].append(nombre)

	return agrupadas


def generar_csv_logs(logs):
	si = io.StringIO()
	cw = csv.writer(si)
	cw.writerow(['Usuario', 'Acción', 'Detalles', 'Fecha'])
	cw.writerows(logs)
	return si.getvalue()


def generar_csv_mapa(data, slot_id):
	si = io.StringIO()
	cw = csv.writer(si)
	cw.writerow(['Nombre del Lote', 'Responsable', 'Tarea', 'Estado'])

	if 'features' in data:
		for f in data['features']:
			props = f.get('properties', {})
			nombre = props.get('name', 'Sin nombre')
			resp = props.get('responsable', 'Sin asignar')
			tareas = props.get('tareas', [])

			if not tareas:
				cw.writerow([nombre, resp, 'Sin tareas', '-'])
			else:
				for t in tareas:
					estado_norm = _estado_tarea(t)
					estado = 'Completada' if estado_norm == 'verde' else ('En proceso' if estado_norm == 'amarillo' else 'Pendiente')
					texto = t.get('texto', '')
					cw.writerow([nombre, resp, texto, estado])

	csv_str = si.getvalue()
	nombre_archivo = f"Avances_Proyecto_{slot_id}_{datetime.now().strftime('%Y%m%d')}.csv"
	return csv_str, nombre_archivo


def generar_csv_avance_por_predio(data, slot_id):
	"""Genera un CSV con la 'Tabla de Avance por Predio' similar al reporte Word.

	Columnas: Lote, Hectareas, Tareas Asignadas, Tareas en Proceso, Tareas Completadas,
	Tareas No Iniciadas, Inversion, Observaciones
	"""
	si = io.StringIO()
	cw = csv.writer(si)
	cw.writerow(['Lote', 'Hectareas', 'Tareas Asignadas', 'Tareas en Proceso', 'Tareas Completadas', 'Tareas No Iniciadas', 'Inversion', 'Observaciones'])

	features = data.get('features', []) if isinstance(data, dict) else []
	predios = {}

	for f in features:
		props = f.get('properties', {}) or {}
		nombre_lote = str(props.get('name', 'Sin nombre') or 'Sin nombre')
		nombre_predio = str(props.get('parent') or nombre_lote)
		hectareas = _a_float(props.get('hectareas', 0))
		inversion = _a_float(props.get('costo', 0))
		observaciones = str(props.get('observaciones') or props.get('incidencias') or '')

		if nombre_predio not in predios:
			predios[nombre_predio] = []
		predios[nombre_predio].append({
			'nombre_lote': nombre_lote,
			'hectareas': hectareas,
			'inversion': inversion,
			'observaciones': observaciones,
			'tareas': props.get('tareas', []) or []
		})

	# Escribir filas por cada lote (predio / lote)
	for nombre_predio, lotes in predios.items():
		for lote in lotes:
			tareas = lote['tareas']
			tareas_asignadas = len(tareas)
			tareas_completadas = sum(1 for t in tareas if _estado_tarea(t) == 'verde')
			tareas_en_proceso = sum(1 for t in tareas if _estado_tarea(t) == 'amarillo')
			tareas_no_iniciadas = tareas_asignadas - tareas_completadas - tareas_en_proceso

			cw.writerow([
				f"{nombre_predio} / {lote['nombre_lote']}",
				f"{lote['hectareas']:.2f}",
				str(tareas_asignadas),
				str(tareas_en_proceso),
				str(tareas_completadas),
				str(tareas_no_iniciadas),
				f"{lote['inversion']:.2f}",
				lote['observaciones']
			])

	csv_str = si.getvalue()
	nombre_archivo = f"Tabla_Avance_Predio_{slot_id}_{datetime.now().strftime('%Y%m%d')}.csv"
	return csv_str, nombre_archivo


def generar_excel_avance_por_predio(data, slot_id):
	wb = Workbook()
	ws_resumen = wb.active
	ws_resumen.title = 'Resumen Ejecutivo'
	ws_avance = wb.create_sheet('Avance por Predio')
	ws_desglose = wb.create_sheet('Desglose por Predio')

	# Estilos base
	fill_header = PatternFill('solid', fgColor='2F5597')
	font_header = Font(color='FFFFFF', bold=True)
	fill_title = PatternFill('solid', fgColor='D9EAF7')
	border = Border(
		left=Side(style='thin', color='D9E2F2'),
		right=Side(style='thin', color='D9E2F2'),
		top=Side(style='thin', color='D9E2F2'),
		bottom=Side(style='thin', color='D9E2F2')
	)
	center = Alignment(horizontal='center', vertical='center', wrap_text=True)
	left = Alignment(horizontal='left', vertical='center', wrap_text=True)

	features = data.get('features', []) if isinstance(data, dict) else []
	predios = {}
	total_tareas = 0
	completadas = 0

	for f in features:
		props = f.get('properties', {}) or {}
		nombre_lote = str(props.get('name', 'Sin nombre') or 'Sin nombre')
		nombre_predio = str(props.get('parent') or nombre_lote)
		hectareas = _a_float(props.get('hectareas', 0))
		inversion = _a_float(props.get('costo', 0))
		observaciones = str(props.get('observaciones') or props.get('incidencias') or '')
		tareas = props.get('tareas', []) or []

		if nombre_predio not in predios:
			predios[nombre_predio] = []
		predios[nombre_predio].append({
			'nombre_lote': nombre_lote,
			'hectareas': hectareas,
			'inversion': inversion,
			'observaciones': observaciones,
			'tareas': tareas
		})

		total_tareas += len(tareas)
		completadas += sum(1 for t in tareas if _estado_tarea(t) == 'verde')

	avance_global = (completadas / total_tareas * 100) if total_tareas > 0 else 0

	# Hoja resumen
	ws_resumen.merge_cells('A1:B1')
	ws_resumen['A1'] = 'Reporte de Avances de Predios - Adhesa'
	ws_resumen['A1'].fill = fill_title
	ws_resumen['A1'].font = Font(bold=True, size=14)
	ws_resumen['A1'].alignment = left
	resumen_rows = [
		('Fecha de Reporte', datetime.now().strftime('%d/%m/%Y %H:%M')),
		('Avance Global', f'{avance_global:.2f}%'),
		('Avance Total', f'{completadas} de {total_tareas} actividades completadas')
	]
	for row_idx, (label, value) in enumerate(resumen_rows, start=3):
		ws_resumen[f'A{row_idx}'] = label
		ws_resumen[f'B{row_idx}'] = value
		ws_resumen[f'A{row_idx}'].font = Font(bold=True)
		ws_resumen[f'A{row_idx}'].alignment = left
		ws_resumen[f'B{row_idx}'].alignment = left

	ws_resumen.column_dimensions['A'].width = 24
	ws_resumen.column_dimensions['B'].width = 48

	# Hoja avance por predio
	headers = ['Lote', 'Hectareas', 'Tareas Asignadas', 'Tareas en Proceso', 'Tareas Completadas', 'Tareas No Iniciadas', 'Inversion', 'Observaciones']
	ws_avance.append(headers)
	for cell in ws_avance[1]:
		cell.fill = fill_header
		cell.font = font_header
		cell.alignment = center
		cell.border = border

	for nombre_predio, lotes in predios.items():
		for lote in lotes:
			tareas = lote['tareas']
			tareas_asignadas = len(tareas)
			tareas_completadas = sum(1 for t in tareas if _estado_tarea(t) == 'verde')
			tareas_en_proceso = sum(1 for t in tareas if _estado_tarea(t) == 'amarillo')
			tareas_no_iniciadas = tareas_asignadas - tareas_completadas - tareas_en_proceso
			ws_avance.append([
				f"{nombre_predio} / {lote['nombre_lote']}",
				lote['hectareas'],
				tareas_asignadas,
				tareas_en_proceso,
				tareas_completadas,
				tareas_no_iniciadas,
				lote['inversion'],
				lote['observaciones']
			])

	for row in ws_avance.iter_rows(min_row=2):
		for cell in row:
			cell.border = border
			cell.alignment = left if cell.column in (1, 8) else center
		if row[1].row % 2 == 0:
			for cell in row:
				cell.fill = PatternFill('solid', fgColor='F8FBFF')

	# Hoja desglose
	desglose_headers = ['Predio', 'Lote', 'Actividad', 'Estado', 'Fecha Inicio', 'Fecha Fin']
	ws_desglose.append(desglose_headers)
	for cell in ws_desglose[1]:
		cell.fill = fill_header
		cell.font = font_header
		cell.alignment = center
		cell.border = border

	for nombre_predio, lotes in predios.items():
		for lote in lotes:
			tareas = lote['tareas']
			if not tareas:
				ws_desglose.append([nombre_predio, lote['nombre_lote'], 'Sin tareas', '-', '-', '-'])
			else:
				for tarea in tareas:
					estado_norm = _estado_tarea(tarea)
					estado = 'Completada' if estado_norm == 'verde' else ('En proceso' if estado_norm == 'amarillo' else 'Pendiente')
					ws_desglose.append([
						nombre_predio,
						lote['nombre_lote'],
						str(tarea.get('texto', 'Sin descripcion') or 'Sin descripcion'),
						estado,
						str(tarea.get('fecha_inicio') or ''),
						str(tarea.get('fecha_fin') or '')
					])

	for row in ws_desglose.iter_rows(min_row=2):
		for cell in row:
			cell.border = border
			cell.alignment = left
		if row[0].row % 2 == 0:
			for cell in row:
				cell.fill = PatternFill('solid', fgColor='F8FBFF')

	for ws in (ws_avance, ws_desglose):
		ws.freeze_panes = 'A2'
		ws.auto_filter.ref = ws.dimensions

	for ws in (ws_avance, ws_desglose):
		for column_cells in ws.columns:
			max_length = 0
			column = column_cells[0].column
			for cell in column_cells:
				try:
					value = str(cell.value) if cell.value is not None else ''
					if len(value) > max_length:
						max_length = len(value)
				except Exception:
					pass
			ws.column_dimensions[get_column_letter(column)].width = min(max_length + 4, 42)

	output = io.BytesIO()
	wb.save(output)
	output.seek(0)
	nombre_archivo = f"Tabla_Avance_Predio_{slot_id}_{datetime.now().strftime('%Y%m%d')}.xlsx"
	return output, nombre_archivo


def generar_word_mapa(data, slot_id):
	doc = Document()
	
	# Titulo principal
	titulo = doc.add_heading('Reporte de Avances de Predios - Adhesa', 0)
	titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
	
	# ===== 1. RESUMEN EJECUTIVO =====
	doc.add_heading('1. Resumen Ejecutivo', level=1)
	
	fecha_reporte = datetime.now().strftime('%d/%m/%Y %H:%M')
	total_tareas = 0
	completadas = 0

	features = data.get('features', [])
	predios = {}
	
	for f in features:
		props = f.get('properties', {})
		tareas = props.get('tareas', [])
		nombre_lote = str(props.get('name', 'Sin nombre') or 'Sin nombre')
		nombre_predio = str(props.get('parent') or nombre_lote)
		hectareas = _a_float(props.get('hectareas', 0))
		inversion = _a_float(props.get('costo', 0))
		observaciones = str(props.get('observaciones') or props.get('incidencias') or 'Sin observaciones')

		if nombre_predio not in predios:
			predios[nombre_predio] = []
		predios[nombre_predio].append({
			'nombre_lote': nombre_lote,
			'hectareas': hectareas,
			'inversion': inversion,
			'observaciones': observaciones,
			'tareas': tareas
		})
		
		total_tareas += len(tareas)
		completadas += sum(1 for t in tareas if _estado_tarea(t) == 'verde')
	
	avance_global = (completadas / total_tareas * 100) if total_tareas > 0 else 0
	
	doc.add_paragraph(f"Fecha de Reporte: {fecha_reporte}")
	doc.add_paragraph(f"Avance Global: {avance_global:.2f}%")
	doc.add_paragraph(f"Avance Total: {completadas} de {total_tareas} actividades completadas")
	
	# ===== 2. TABLA DE AVANCE POR PREDIO =====
	doc.add_heading('2. Tabla de Avance por Predio', level=1)
	
	# Crear tabla con columnas solicitadas
	table = doc.add_table(rows=1, cols=8)
	table.style = 'Table Grid'
	
	hdr_cells = table.rows[0].cells
	hdr_cells[0].text = 'Lote'
	hdr_cells[1].text = 'Hectareas'
	hdr_cells[2].text = 'Tareas Asignadas'
	hdr_cells[3].text = 'Tareas en Proceso'
	hdr_cells[4].text = 'Tareas Completadas'
	hdr_cells[5].text = 'Tareas No Iniciadas'
	hdr_cells[6].text = 'Inversion'
	hdr_cells[7].text = 'Observaciones'
	
	for hdr_cell in hdr_cells:
		hdr_cell._element.get_or_add_tcPr().append(
			doc.element.makeelement('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
		)
	
	for nombre_predio, lotes in predios.items():
		for lote in lotes:
			tareas = lote['tareas']
			tareas_asignadas = len(tareas)
			tareas_completadas = sum(1 for t in tareas if _estado_tarea(t) == 'verde')
			tareas_en_proceso = sum(1 for t in tareas if _estado_tarea(t) == 'amarillo')
			tareas_no_iniciadas = tareas_asignadas - tareas_completadas - tareas_en_proceso

			row_cells = table.add_row().cells
			row_cells[0].text = f"{nombre_predio} / {lote['nombre_lote']}"
			row_cells[1].text = f"{lote['hectareas']:.2f}"
			row_cells[2].text = str(tareas_asignadas)
			row_cells[3].text = str(tareas_en_proceso)
			row_cells[4].text = str(tareas_completadas)
			row_cells[5].text = str(tareas_no_iniciadas)
			row_cells[6].text = f"${lote['inversion']:,.2f}"
			row_cells[7].text = lote['observaciones']
	
	# ===== 3. DESGLOSE POR PREDIO =====
	doc.add_heading('3. Desglose por Predio', level=1)

	for nombre_predio, lotes in predios.items():
		doc.add_heading(f"Predio: {nombre_predio}", level=2)

		for lote in lotes:
			nombre_lote = lote['nombre_lote']
			hectareas = lote['hectareas']
			tareas = lote['tareas']

			doc.add_paragraph(f"Nombre del lote: {nombre_lote}")
			doc.add_paragraph(f"Hectareas: {hectareas:.2f} ha")

			agrupadas = _agrupar_actividades_por_estado(tareas)
			doc.add_paragraph(
				"Actividades asignadas: "
				f"Completadas ({len(agrupadas['completadas'])}): {', '.join(agrupadas['completadas']) if agrupadas['completadas'] else 'Ninguna'} | "
				f"En proceso ({len(agrupadas['en_proceso'])}): {', '.join(agrupadas['en_proceso']) if agrupadas['en_proceso'] else 'Ninguna'} | "
				f"No iniciadas ({len(agrupadas['no_iniciadas'])}): {', '.join(agrupadas['no_iniciadas']) if agrupadas['no_iniciadas'] else 'Ninguna'}"
			)

			fechas_inicio = [t.get('fecha_inicio') for t in tareas if t.get('fecha_inicio')]
			fechas_fin = [t.get('fecha_fin') for t in tareas if t.get('fecha_fin')]
			fecha_inicio_min = min(fechas_inicio) if fechas_inicio else 'N/A'
			fecha_fin_max = max(fechas_fin) if fechas_fin else 'N/A'

			doc.add_paragraph(f"Inicio de actividades: {fecha_inicio_min}")
			doc.add_paragraph(f"Finalizacion de actividades: {fecha_fin_max}")
			doc.add_paragraph()
	
	target = io.BytesIO()
	doc.save(target)
	target.seek(0)
	return target, f"Reporte_Avances_Slot_{slot_id}.docx"
